"""Step-by-step User Guide captures for the 911 Setup family.

Dev tool, run from the repo root (one app per process):

    python -u tools/guide_captures/cap_911_setup_family.py check    # fixture self-test, no GUI
    python -u tools/guide_captures/cap_911_setup_family.py setup    # 911 Setup, full flow
    python -u tools/guide_captures/cap_911_setup_family.py cards    # 911 Teams Cards picker
    python -u tools/guide_captures/cap_911_setup_family.py ticket   # 911 Remove Ticket prompt

Uses the camera library (tools/capture_guide_screens.py): sandboxed fresh
profile, hidden main window. App-created dialogs flash on the real screen
briefly; expected for a manual capture tool.

Everything modal is driven from a QTimer state machine (probe()'s pattern) --
a modal dialog's nested event loop freezes straight-line code, but timers
keep firing inside it.

FIXTURES are all fake practice data (batch V060, invented nests/parts/orders)
and live under C:\\Users\\Public\\Pilot Program -- NOT the pid-stamped temp
sandbox -- because the app logs its QTDR root to the console and a temp path
(with the capturing account's name in it) must never appear in a picture.
Delete that folder when the captures are done.
"""

from __future__ import annotations

import datetime as dt
import os
import shutil
import sys
import time
import traceback
from pathlib import Path

sys.path.insert(0, r"C:\Dev\TechDeck\tools")
import capture_guide_screens as cam  # noqa: E402  (sandboxes LOCALAPPDATA on import)

REPO = Path(r"C:\Dev\TechDeck")

# ── Fixture layout (neutral, believable paths -- they show in console logs) ──
PUBLIC_ROOT = Path(r"C:\Users\Public\Pilot Program")
QTDR = PUBLIC_ROOT / "911 QTDR"
TEMPLATE_DIR = QTDR / "03 - Processing Forms & Templates" / "00 - SACO"
FORECAST_DIR = PUBLIC_ROOT / "Forecast and Inventory Reports"
SCHEDULE = (PUBLIC_ROOT / "922 QTDR Production Packages" / "2 - Planning"
            / "EB 922 Schedule.xlsx")

BATCH = "V060"

# nest -> [(work order, DYPN, stock code, qty, scope)]
NESTS = {
    "504101": [
        ("DL482911", "BK573423-11", "211078455", 3, "CUT TO LENGTH"),
        ("DL482912", "BK573424-12", "211078455", 2, "CUT TO LENGTH"),
    ],
    "504102": [
        ("DL482921", "BK573431-21", "211080122", 4, "CUT TO LENGTH"),
        ("DL482922", "BK573432-22", "211080122", 1, "CUT TO LENGTH"),
    ],
}
NEST_MATERIAL = {"504101": "HSS 4 X 4 X 3/8", "504102": "FLAT BAR 3 X 1/2"}
NEST_SIZE = {"504101": "4 X 4 X 3/8", "504102": "3 X 1/2"}
MIL_SPEC = "MIL-S-22698"

# Schedule RATING legend fills (literal RGB, from the difficulty reader)
FILL_SIMPLE, FILL_MEDIUM, FILL_DIFFICULT = "DAF2D0", "F9EC8F", "FA949B"


# ─────────────────────────────────────────────────────────────────────────────
# Fixture builders
# ─────────────────────────────────────────────────────────────────────────────

def _packet_pdf(path: Path, nest: str, parts, material: str, size_text: str):
    """A fake nest packet: cover (with the stamp anchors), SUMMARY OF NEST,
    one MOVE TICKET page, one PART SKETCH page. Landscape letter."""
    import fitz

    doc = fitz.open()
    W, H = 792, 612

    # -- page 1: cover ------------------------------------------------------
    p = doc.new_page(width=W, height=H)
    p.insert_text((60, 55), "911 NEST PACKAGE", fontsize=22, fontname="hebo")
    p.insert_text((60, 90), f"NEST PACKAGE NUMBER: {nest}", fontsize=12)
    p.insert_text((60, 115), "ORDER TYPE: QTDR", fontsize=10)
    p.insert_text((60, 133), "SHIP VIA: SUPPLIER TRUCK", fontsize=10)
    p.draw_line((55, 150), (740, 150), width=0.8)

    # Quality Requirements grid; the batch/nest stamp lands ~67pt below it.
    p.insert_text((430, 230), "Quality Requirements", fontsize=12, fontname="hebo")
    p.draw_rect(fitz.Rect(428, 240, 700, 294), width=0.8)
    p.draw_line((428, 258), (700, 258), width=0.5)
    p.draw_line((428, 276), (700, 276), width=0.5)

    # Bottom form row: Material Type / Material Size / Alternate Source Code.
    p.insert_text((60, 480), "Material Type", fontsize=11, fontname="hebo")
    p.insert_text((240, 480), "Material Size", fontsize=11, fontname="hebo")
    p.insert_text((430, 480), "Alternate Source Code", fontsize=11, fontname="hebo")
    p.insert_text((240, 502), size_text, fontsize=10)
    p.draw_rect(fitz.Rect(55, 486, 175, 516), width=0.6)
    p.draw_rect(fitz.Rect(235, 486, 355, 516), width=0.6)
    p.draw_rect(fitz.Rect(410, 490, 580, 532), width=0.6)

    # -- page 2: SUMMARY OF NEST (flat line sequence the parser expects) ----
    p = doc.new_page(width=W, height=H)
    p.insert_text((60, 55), "SUMMARY OF NEST", fontsize=16, fontname="hebo")
    y = 95
    for line in ("REF", "PART NUMBER", "QTY", "WORK ORDER"):
        p.insert_text((60, y), line, fontsize=10, fontname="hebo")
        y += 14
    y += 6
    for i, (wo, dypn, _code, qty, _scope) in enumerate(parts, 1):
        for line in (str(i), dypn, str(qty), wo):
            p.insert_text((60, y), line, fontsize=10)
            y += 14

    # -- page 3: MOVE TICKET (removed by the omit build; feeds MIL/MATERIAL) --
    p = doc.new_page(width=W, height=H)
    p.insert_text((60, 55), "MOVE TICKET", fontsize=18, fontname="hebo")
    wo, dypn, _code, qty, _scope = parts[0]
    y = 100
    for line in (f"WORK ORDER: {wo}", f"DYPN: {dypn}", f"QTY: {qty}",
                 f"MIL SPEC: {MIL_SPEC}", f"MATERIAL: {material}", "LEVEL: N"):
        p.insert_text((60, y), line, fontsize=11)
        y += 20

    # -- page 4: PART SKETCH (kept) -----------------------------------------
    p = doc.new_page(width=W, height=H)
    p.insert_text((60, 55), "PART SKETCH", fontsize=18, fontname="hebo")
    p.insert_text((60, 85), f"DYPN: {parts[0][1]}", fontsize=11)
    p.draw_rect(fitz.Rect(200, 150, 600, 400), width=1.2)
    p.draw_circle(fitz.Point(400, 275), 60, width=1.0)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def _batch_list(path: Path, batch: str, nests: dict):
    """A BATCH LIST workbook: 'BATCH' sheet, headers in row 3, data row 4+."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BATCH"
    ws["A1"] = f"({batch}) BATCH LIST"
    headers = ["Work Order", "DYPN", "Material", "DYPN QTY",
               "Material Amount (Total)", "Nest Pkg Nbr", "SCOPE OF WORK",
               "Description"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c).value = h
    r = 4
    for nest, parts in nests.items():
        for (wo, dypn, code, qty, scope) in parts:
            desc = NEST_MATERIAL.get(nest, "FLAT BAR")
            ws.cell(r, 1).value = wo
            ws.cell(r, 2).value = dypn
            ws.cell(r, 3).value = code
            ws.cell(r, 4).value = qty
            ws.cell(r, 5).value = round(qty * 3.7, 1)  # deliberately NOT the qty
            ws.cell(r, 6).value = nest
            ws.cell(r, 7).value = scope
            ws.cell(r, 8).value = desc
            r += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _forecast(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "911 Forecast"
    for c, h in enumerate(["PO", "Line", "Batch /DR", "Nest", "TRACE/MIC"], 1):
        ws.cell(1, c).value = h
    ws.append(["4500178231", 10, BATCH, "504101", "DL34270"])
    ws.append(["4500178231", 20, BATCH, "504102", "XL30183"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _template_workbook(path: Path):
    """The '911 BATCH _.xlsx' template: NEST + SCRIBE VERIFICATION sheets."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NEST"
    ws["A1"] = "911 BATCH WORKBOOK"
    headers = ["PO", "LINE", "BATCH /DR", "MIL SPEC", "MATL TYPE",
               "WORK ORDER", "DYPN", "MATERIAL", "DYPN QTY", "NEST PKG NBR",
               "SCOPE OF WORK"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c).value = h
    scr = wb.create_sheet("SCRIBE VERIFICATION")
    for c, h in enumerate(["QTY", "DYPN", "HULL CODE", "MILL - SPEC",
                           "MAT TYPE", "UNIQUE - TRACE"], 1):
        scr.cell(1, c).value = h
    wb.create_sheet("INSPECTION SHEET")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _schedule(path: Path):
    """EB 922 Schedule: CURRENT PIPELINE with RATING carried by fill colour."""
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CURRENT PIPELINE"
    for c, h in enumerate(["DEPT.", "BATCH / NEST", "DATE", "NOTES", "RATING",
                           "STATUS"], 1):
        ws.cell(1, c).value = h

    rows = [
        # the V060 nests 911 Setup stamps (already carded + set up queue-wise)
        ("911", "V060 504101", dt.datetime(2026, 9, 18), "HSS 4 X 4 X 3/8",
         FILL_MEDIUM, "NEED MODEL"),
        ("911", "V060 504102", dt.datetime(2026, 9, 18), "FLAT BAR 3 X 1/2",
         FILL_SIMPLE, "NEED MODEL"),
        # the Teams Cards queue
        ("911", "V102 504098", dt.datetime(2026, 10, 30),
         "TUBE 2.5 OD X .25 WALL", FILL_MEDIUM, "NEED TEAMS/SETUP"),
        ("911", "V102 504099", dt.datetime(2026, 11, 6),
         "ANGLE 3 X 3 X 3/8", FILL_SIMPLE, "NEED TEAMS/SETUP"),
        ("911", "V103 505210", "HOLD",
         "FLAT BAR 4 X 1/2 (TL)", FILL_DIFFICULT, "NEED TEAMS/SETUP"),
    ]
    for r, (dept, key, date, notes, fill, status) in enumerate(rows, 2):
        ws.cell(r, 1).value = dept
        ws.cell(r, 2).value = key
        ws.cell(r, 3).value = date
        ws.cell(r, 4).value = notes
        ws.cell(r, 5).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, 6).value = status
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def build_fixtures():
    if PUBLIC_ROOT.exists():
        shutil.rmtree(PUBLIC_ROOT, ignore_errors=True)

    # 911 Setup's batch
    _batch_list(QTDR / BATCH / f"{BATCH} BATCH LIST.xlsx", BATCH, NESTS)
    for nest, parts in NESTS.items():
        _packet_pdf(QTDR / BATCH / "NEST PACKAGES" / f"{nest}.pdf",
                    nest, parts, NEST_MATERIAL[nest], NEST_SIZE[nest])
    _template_workbook(TEMPLATE_DIR / "911 BATCH _.xlsx")
    (TEMPLATE_DIR / "QF-QU-15 REV B - SCRIBE VERIFICATION - SHAPES.docx"
     ).write_bytes(b"practice scribe form placeholder")
    _forecast(FORECAST_DIR / "Working Forecast List.xlsx")
    _schedule(SCHEDULE)

    # Teams Cards' referenced batches (stock codes for the card titles)
    _batch_list(QTDR / "V102" / "V102 BATCH LIST.xlsx", "V102", {
        "504098": [("DL491210", "BK574810-10", "211078455", 2, "CUT TO LENGTH")],
        "504099": [("DL491211", "BK574811-11", "211080122", 5, "CUT TO LENGTH")],
    })
    _batch_list(QTDR / "V103" / "V103 BATCH LIST.xlsx", "V103", {
        "505210": [("DL491305", "BK575105-05", "211055890", 3, "CUT TO LENGTH")],
    })
    print(f"Fixtures built under {PUBLIC_ROOT}")


# ─────────────────────────────────────────────────────────────────────────────
# Fixture self-test (no GUI): run the real parsers over the fixtures
# ─────────────────────────────────────────────────────────────────────────────

def _load_plugin(pid: str):
    import importlib.util
    p = REPO / "plugins" / pid / "run.py"
    spec = importlib.util.spec_from_file_location(f"chk_{pid}", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def check():
    build_fixtures()
    import fitz

    su = _load_plugin("911_setup")
    rt = _load_plugin("911_remove_ticket")
    tc = _load_plugin("911_teams_cards")

    pdf = QTDR / BATCH / "NEST PACKAGES" / "504101.pdf"
    q = su._parse_packet_summary_qtys(pdf)
    assert q == {("BK573423-11", "DL482911"): 3,
                 ("BK573424-12", "DL482912"): 2}, q

    mil, mat, fer = su._extract_pdf_data(pdf)
    assert mil == MIL_SPEC and mat == NEST_MATERIAL["504101"], (mil, mat, fer)

    doc = fitz.open(str(pdf))
    remove, mats = rt._scan_document(doc)
    assert remove == {2} and mats == {NEST_MATERIAL["504101"]}, (remove, mats)
    words = doc[0].get_text("words")
    assert rt._find_header(words, "Quality", "Requirements")
    assert rt._find_header(words, "Material", "Type")
    assert rt._find_header_seq(words, "Alternate", "Source", "Code")
    doc.close()

    dm, prob = rt._load_difficulty_map(
        {"settings": {"schedule_path": str(SCHEDULE)}}, print)
    assert not prob, prob
    assert dm.get("504101") == "MEDIUM" and dm.get("504102") == "SIMPLE", dm
    assert dm.get("505210") == "DIFFICULT" and dm.get("504098") == "MEDIUM", dm

    mapping, warn = tc._read_batch_list_materials(QTDR, "V102", print)
    assert mapping.get("504098", {}).get("code") == "211078455", (mapping, warn)

    nests = su._get_unique_nests_from_batch_list(
        QTDR / BATCH / f"{BATCH} BATCH LIST.xlsx")
    assert nests == ["504101", "504102"], nests

    print("CHECK OK")


# ─────────────────────────────────────────────────────────────────────────────
# Capture runs (QTimer state machines, probe()'s event-driven pattern)
# ─────────────────────────────────────────────────────────────────────────────

def _arm_safety(seconds: int = 90):
    """Global watchdog: whatever happens, the process ends."""
    from PySide6.QtCore import QTimer
    t = QTimer()
    t.setSingleShot(True)

    def bail():
        print(f"SAFETY TIMEOUT after {seconds}s -- aborting capture.")
        sys.stdout.flush()
        os._exit(1)

    t.timeout.connect(bail)
    t.start(seconds * 1000)
    return t


def _find_dialog(app_window, cls_name: str, title_sub: str | None = None):
    from PySide6.QtWidgets import QApplication
    for w in QApplication.topLevelWidgets():
        if w is app_window or not w.isVisible():
            continue
        if type(w).__name__ != cls_name:
            continue
        if title_sub and title_sub.lower() not in (w.windowTitle() or "").lower():
            continue
        return w
    return None


def _copy_shot(src_name: str, dest_name: str):
    shutil.copyfile(cam.OUT_DIR / f"{src_name}.png",
                    cam.OUT_DIR / f"{dest_name}.png")
    print(f"  {dest_name}.png  (copy of {src_name}.png)")


def run_setup():
    """911 Setup: actions window -> batch prompt -> nest picker -> QTY verify
    -> finished console. Inspection Sheets stays unchecked (it drives the real
    Excel via COM, which a capture run must not do)."""
    build_fixtures()
    app, window, settings = cam.boot()
    settings.set_plugin_settings("911_setup", {
        "qtdr_base_path": str(QTDR),
        "forecast_dir": str(FORECAST_DIR),
        "schedule_path": str(SCHEDULE),
        "card_dry_run": True,   # belt and braces; the cards stage stays off
    })

    from PySide6.QtCore import Qt, QTimer

    _safety = _arm_safety(90)  # noqa: F841  (kept alive by reference)
    st = {"phase": "actions", "stable": 0}
    timer_ref = {}

    def step():
        con = window.console
        phase = st["phase"]

        if phase == "actions":
            dlg = _find_dialog(window, "GroupedToggleDialog")
            if dlg is not None:
                st["stable"] += 1
                if st["stable"] >= 3:
                    cam.save(dlg, "911_setup_actions")
                    dlg._parents["inspection_sheets"].setCheckState(
                        0, Qt.CheckState.Unchecked)
                    st.update(phase="actions_go", stable=0)

        elif phase == "actions_go":
            dlg = _find_dialog(window, "GroupedToggleDialog")
            if dlg is not None:
                dlg.run_btn.click()
                st.update(phase="batch", stable=0)

        elif phase == "batch":
            prompt = (getattr(con, "input_prompt", "") or "").lower()
            if con.waiting_for_input and "batch number" in prompt:
                st["stable"] += 1
                if st["stable"] >= 2:
                    cam.save(con, "911_setup_batch_prompt")
                    _copy_shot("911_setup_batch_prompt", "run_console_prompt")
                    con.input_field.setText(BATCH)
                    con._on_input_submitted()
                    st.update(phase="nest", stable=0)

        elif phase == "nest":
            dlg = _find_dialog(window, "NestSelectionDialog")
            if dlg is not None:
                st["stable"] += 1
                if st["stable"] == 2:
                    dlg.root.setCheckState(0, Qt.CheckState.Checked)
                if st["stable"] >= 4:
                    cam.save(dlg, "911_setup_nest_selection")
                    # The QTY verification lines print (and scroll past) fast;
                    # tick quickly until that shot is in the can.
                    timer_ref["t"].setInterval(50)
                    dlg.run_btn.click()
                    st.update(phase="qty", stable=0)

        elif phase == "qty":
            # The worker outruns the console: its log lines sit in
            # RunController's buffer and a 50ms timer delivers them in batches
            # of 15, so "the moment the verification lines are the newest
            # thing on screen" can be skipped entirely. Take over the drain:
            # stop the timer, deliver the buffered lines one at a time through
            # the same signal, and grab the instant the 'Verified' line lands.
            import queue as _q
            rc = window.home_page._run
            if not st.get("drain_stopped"):
                rc._log_drain_timer.stop()
                st["drain_stopped"] = True
            done_shot = False
            while True:
                try:
                    tid, msg = rc._log_buffer.get_nowait()
                except _q.Empty:
                    break
                window.home_page.plugin_log.emit(tid, msg)
                if "Verified: the column labeled" in msg:
                    cam.save(con, "911_setup_qty_verify")
                    done_shot = True
                    break
            if done_shot or "911 Setup complete" in con.output.toPlainText():
                rc._log_drain_timer.start()
                timer_ref["t"].setInterval(220)
                st.update(phase="done", stable=0)

        elif phase == "done":
            if "911 Setup complete for batch" in con.output.toPlainText():
                st["stable"] += 1
                if st["stable"] >= 5:
                    cam.save(con, "911_setup_run_complete")
                    print("SETUP CAPTURE DONE")
                    sys.stdout.flush()
                    os._exit(0)

    def tick():
        try:
            step()
        except Exception:
            traceback.print_exc()
            sys.stdout.flush()
            os._exit(2)

    timer = QTimer()
    timer_ref["t"] = timer
    timer.timeout.connect(tick)
    timer.start(220)
    cam.start_app(app, window, "911_setup")
    app.exec()


def run_cards():
    """911 Teams Cards: the 'Select Nests to Card' picker, then Cancel.
    card_dry_run is forced ON so nothing could ever post, and the picker is
    cancelled anyway -- no webhook call, no schedule write."""
    build_fixtures()
    app, window, settings = cam.boot()
    settings.set_plugin_settings("911_teams_cards", {
        "qtdr_base_path": str(QTDR),
        "schedule_path": str(SCHEDULE),
        "card_dry_run": True,
    })
    # boot()'s everyday kit doesn't include this app; the run queue only
    # takes tiles that are in the kit.
    tiles = settings.get_profile_tiles()
    if "911_teams_cards" not in tiles:
        settings.set_profile_tiles(tiles + ["911_teams_cards"])
    window.home_page._refresh_tiles()
    cam.pump(app, 400)

    from PySide6.QtCore import QTimer

    _safety = _arm_safety(90)  # noqa: F841
    st = {"phase": "picker", "stable": 0}

    def step():
        if st["phase"] == "picker":
            dlg = _find_dialog(window, "SelectionDialog", "911 Teams Cards")
            if dlg is not None:
                st["stable"] += 1
                if st["stable"] == 2:
                    dlg.resize(620, 460)   # min width truncates the due dates
                if st["stable"] >= 4:
                    cam.save(dlg, "911_teams_cards_picker")
                    dlg.reject()   # leave the queue untouched
                    st.update(phase="exit", stable=0)
        elif st["phase"] == "exit":
            st["stable"] += 1
            if st["stable"] >= 5:
                print("CARDS CAPTURE DONE")
                sys.stdout.flush()
                os._exit(0)

    def tick():
        try:
            step()
        except Exception:
            traceback.print_exc()
            sys.stdout.flush()
            os._exit(2)

    timer = QTimer()
    timer.timeout.connect(tick)
    timer.start(220)
    cam.start_app(app, window, "911_teams_cards")
    app.exec()


def run_ticket():
    """911 Remove Ticket: the console at its first typed prompt (the PDF
    folder path). Nothing is answered; the process hard-exits after the shot."""
    build_fixtures()
    app, window, settings = cam.boot()

    from PySide6.QtCore import QTimer

    _safety = _arm_safety(60)  # noqa: F841
    st = {"stable": 0}

    def step():
        con = window.console
        prompt = (getattr(con, "input_prompt", "") or "").lower()
        if con.waiting_for_input and "pdf directory" in prompt:
            st["stable"] += 1
            if st["stable"] >= 2:
                cam.save(con, "911_remove_ticket_prompt")
                print("TICKET CAPTURE DONE")
                sys.stdout.flush()
                os._exit(0)

    def tick():
        try:
            step()
        except Exception:
            traceback.print_exc()
            sys.stdout.flush()
            os._exit(2)

    timer = QTimer()
    timer.timeout.connect(tick)
    timer.start(220)
    cam.start_app(app, window, "911_remove_ticket")
    app.exec()


def main():
    mode = sys.argv[1] if len(sys.argv) > 1 else ""
    if mode == "check":
        check()
    elif mode == "setup":
        run_setup()
    elif mode == "cards":
        run_cards()
    elif mode == "ticket":
        run_ticket()
    else:
        sys.exit("usage: cap_911_setup_family.py check|setup|cards|ticket")
    sys.stdout.flush()
    os._exit(0)


if __name__ == "__main__":
    main()
