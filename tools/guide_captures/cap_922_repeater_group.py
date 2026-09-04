"""User Guide captures: 922 Batch Repeater (full step-by-step), plus
922 Kitting, 922 LST Organizer and 922 Runtime Genie (1-2 pictures each).

Run from the repo root, one scene per invocation (each boots a fresh
sandbox profile):

    PYTHONUTF8=1 python -u tools/guide_captures/cap_922_repeater_group.py --scene repeater
    ... --scene kitting | lst | genie          (or --scene all)

Everything runs against the camera library's fresh sandbox profile with a
FAKE 922 tree (Batch 573 plus prior Batches 571/572, invented order numbers)
mounted on a subst'd drive (T: preferred, U:/V: fallback) so no real path,
batch, or user name can appear in a picture. Same fictional batch/orders as
cap_922_setup_family.py so the guide's 922 pictures tell one story.

The 922 Batch Repeater keeps its TYPED batch prompt (typing a new number is
how it creates a batch folder), so its first picture needs no picker at all.
The other three apps pick the `Batch NNN` folder; the Sentry Drone gadget is
unlocked and switched on for them so the pick opens as a drivable Qt window
(the chopper picker) instead of the native Explorer dialog - a native dialog
can neither be photographed nor driven and would hang the run.

All modal windows (the Select Repeats to Pull window, the drone picker, the
LST results popup) are driven from a QTimer state machine: a modal exec()
nests the event loop, so straight-line code freezes until the dialog closes,
but timers keep firing inside the nested loop. A global QTimer safety
hard-exits after 120 s no matter what (plus a threading.Timer backstop in
case the GUI thread is wedged in a loop Qt timers can't reach).

No network: the repeater's Teams webhook post is monkeypatched to a local
no-op success and the usage-telemetry webhook URL is blanked, so a capture
run can never touch the real Planner flow or the usage table.
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
import threading
import time
from pathlib import Path

sys.path.insert(0, r"C:\Dev\TechDeck\tools")
import capture_guide_screens as cam  # noqa: E402  (sets the sandbox FIRST)

BATCH = "573"
ROOT_NAME = "922 QTDR Production Packages"
FAKE_PARENT = cam._SANDBOX / "fx_shop"

_MOUNTED: list = []  # drive letter this process subst'd (for cleanup)

# The same fictional Batch 573 universe as cap_922_setup_family.py:
# (new WO, prior batch, prior WO, PPN, tube serial, material description)
PARTS = [
    ("BK573423", "571", "BJ571408", "R7651664-H23", "218002867", "1.5x1.5x0.188 NOM"),
    ("BK573424", "571", "BJ571411", "R7651665-H23", "218003095", "2x2x0.25 NOM"),
    ("BK573425", "572", "BH572106", "R7651666-H23", "218004492", "1.25x1.25x0.12 NOM"),
    ("X6401069", None, None, "R7651667-H23", "218012302", "1x1x0.083 NOM"),
]
# One PO tube line with no .lst on disk -> the LST report/popup shows a miss.
MISSING = ("X6401069", "R7651667-H23-6", "218019939", "3x2x0.188 NOM")

PIECE = {ppn: f"{ppn}-4" for _, _, _, ppn, _, _ in PARTS}
MINUTES = {"R7651664-H23-4": 24.5, "R7651665-H23-4": 18.75,
           "R7651666-H23-4": 31.0, "R7651667-H23-4": 12.25}

# The MPL PO matrix: which PPNs each prior batch already made.
MPL_COLUMNS = {
    "PO 570": ["R7651665-H23"],
    "PO 571": ["R7651664-H23", "R7651665-H23"],
    "PO 572": ["R7651666-H23"],
    "PO 573": [ppn for _, _, _, ppn, _, _ in PARTS],
}

APP_KIT = ["922_batch_repeater", "922_kitting", "922_lst_organizer",
           "922_runtime_genie"]


def _unmount():
    for letter in _MOUNTED:
        subprocess.run(f"subst {letter} /D", shell=True, capture_output=True)


def finish(code: int, why: str = ""):
    if why:
        print(f"-- exit({code}): {why}", flush=True)
    _unmount()
    sys.stdout.flush()
    os._exit(code)


def mount_drive() -> Path:
    """subst the fake shop onto a free drive letter. Never tears down an
    existing mapping (a sibling capture script may own it)."""
    for letter in ("T:", "U:", "V:"):
        if Path(letter + "\\").exists():
            continue  # in use - leave it alone
        subprocess.run(f'subst {letter} "{FAKE_PARENT}"', shell=True,
                       capture_output=True, text=True)
        root = Path(f"{letter}\\{ROOT_NAME}")
        if root.is_dir():
            _MOUNTED.append(letter)
            print(f"Fake root mounted: {root}")
            return root
        subprocess.run(f"subst {letter} /D", shell=True, capture_output=True)
    finish(1, "no free drive letter for subst")


# --- Fixtures ---------------------------------------------------------------

def _write_lst(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("TechDeck User Guide practice .lst data\n", encoding="ascii")


def _make_pdf(path: Path, title: str, lines):
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 70), title, fontsize=18, fontname="hebo")
    y = 110
    for ln in lines:
        page.insert_text((72, y), ln, fontsize=11)
        y += 20
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def _make_mpl(path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "PO 321+"
    ws["A1"] = "922 MPL"
    for c, (hdr, vals) in enumerate(MPL_COLUMNS.items(), start=1):
        ws.cell(3, c, hdr)
        for r, v in enumerate(vals, start=4):
            ws.cell(r, c, v)
    wb.save(str(path))


def _make_po_workbook(path: Path):
    """The batch's QF-QU-09 workbook: PO sheet + SOURCE MATERIAL sheet."""
    from openpyxl import Workbook
    wb = Workbook()
    po = wb.active
    po.title = "PO"
    po.append(["ORDER", "DYPN", "SOURCE MATERIAL"])
    for wo, _pb, _pw, ppn, serial, _desc in PARTS:
        po.append([wo, PIECE[ppn], serial])
    po.append([MISSING[0], MISSING[1], MISSING[2]])
    sm = wb.create_sheet("SOURCE MATERIAL")
    sm.append(["SOURCE MATERIAL", "PART DESCRIPTION"])
    for _wo, _pb, _pw, _ppn, serial, desc in PARTS:
        sm.append([serial, desc])
    sm.append([MISSING[2], MISSING[3]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _order_folder_with_prints(order_dir: Path, ppn: str, with_pdf: bool):
    seven = order_dir / "CAD-AND-SHOP-PRINTS" / ppn / "7000"
    _write_lst(seven / f"{PIECE[ppn]}.lst")
    if with_pdf:
        _make_pdf(seven / f"{PIECE[ppn]}.pdf", "7000 PROGRAM SHEET",
                  [f"Part: {PIECE[ppn]}",
                   f"Batch: {BATCH}",
                   f"Machine time decimal {MINUTES[PIECE[ppn]]} min"])


def build_fixture(scene: str):
    import shutil
    if FAKE_PARENT.exists():
        shutil.rmtree(FAKE_PARENT)
    root = FAKE_PARENT / ROOT_NAME
    batch = root / f"Batch {BATCH}"
    doc_dir = batch / f"Batch {BATCH} - Documentation"
    for d in (root / "1 - Completed", root / "2 - Planning", doc_dir):
        d.mkdir(parents=True, exist_ok=True)
    _make_mpl(root / "922 MPL.xlsx")

    # Prior batches holding the repeat orders' source folders.
    for _wo, pb, pw, ppn, _serial, _desc in PARTS:
        if pb is None:
            continue
        src = root / f"Batch {pb}" / f"{pw}-{ppn}"
        _order_folder_with_prints(src, ppn, with_pdf=False)
        _make_pdf(src / f"Binder - {pw}.pdf", f"BINDER  {pw}-{ppn}",
                  ["TechDeck User Guide practice data."])

    # The new batch's order folders.
    for wo, _pb, _pw, ppn, _serial, _desc in PARTS:
        od = batch / f"{wo}-{ppn}"
        od.mkdir(parents=True, exist_ok=True)
        if scene in ("lst", "genie"):
            _order_folder_with_prints(od, ppn, with_pdf=(scene == "genie"))

    if scene == "repeater":
        # One repeat already sitting in REPEAT BATCHES from a prior run, so
        # the selection window shows its "already pulled" flag.
        _wo, _pb, pw, ppn, _serial, _desc = PARTS[0]
        pre = batch / "REPEAT BATCHES" / f"{pw}-{ppn}"
        _order_folder_with_prints(pre, ppn, with_pdf=False)

    if scene in ("kitting", "lst", "genie"):
        _make_po_workbook(doc_dir / f"PO H{BATCH} QF-QU-09.xlsx")

    if scene == "genie":
        # As if the 922 LST Organizer already ran: the gathered .lst files
        # filed into their source-material folders under Documentation\LST.
        for _wo, _pb, _pw, ppn, serial, desc in PARTS:
            _write_lst(doc_dir / "LST" / f"{desc} ({serial})"
                       / f"{PIECE[ppn]}.lst")

    print(f"Fixture built for scene '{scene}': {batch}")


# --- Qt driver (runs inside nested modal event loops via QTimer) ------------

class Driver:
    """One QTimer tick loop that answers whatever modal window is up."""

    def __init__(self, app, window):
        self.app = app
        self.window = window
        self.sel = None      # Select Repeats to Pull window
        self.picker = None   # drone folder picker
        self.popup = None    # an expected QMessageBox to photograph
        from PySide6.QtCore import QTimer
        self.timer = QTimer()
        self.timer.timeout.connect(self.tick)
        self.timer.start(150)

    # -- arming --------------------------------------------------------------
    def expect_selection(self, shot: str):
        self.sel = {"step": 0, "t": 0.0, "shot": shot, "done": False}

    def expect_picker(self, shot: str | None, mode: str = "fire"):
        """mode='fire': lock Batch 573, photograph, Execute, skip kill-cam.
        mode='cancel': lock + photograph, then Cancel (no kill-cam)."""
        self.picker = {"step": 0, "t": 0.0, "shot": shot, "mode": mode,
                       "done": False}

    def expect_popup(self, title_sub: str, shot: str):
        self.popup = {"step": 0, "t": 0.0, "title": title_sub.lower(),
                      "shot": shot, "done": False}

    def sel_done(self):
        return self.sel is not None and self.sel["done"]

    def picker_done(self):
        return self.picker is not None and self.picker["done"]

    def popup_done(self):
        return self.popup is not None and self.popup["done"]

    # -- helpers ---------------------------------------------------------------
    def _find(self, cls_name: str, title_sub: str = ""):
        from PySide6.QtWidgets import QApplication
        for w in QApplication.topLevelWidgets():
            if (w.isVisible() and type(w).__name__ == cls_name
                    and title_sub.lower() in (w.windowTitle() or "").lower()):
                return w
        return None

    def _handle_messageboxes(self):
        from PySide6.QtWidgets import QApplication, QMessageBox
        for w in QApplication.topLevelWidgets():
            if not (isinstance(w, QMessageBox) and w.isVisible()):
                continue
            title = (w.windowTitle() or "").lower()
            if (self.popup is not None and not self.popup["done"]
                    and self.popup["title"] in title):
                self._tick_popup(w)
                continue
            print(f"!! unexpected popup: {w.windowTitle()!r}: "
                  f"{w.text()[:300]}", flush=True)
            w.grab().save(str(FAKE_PARENT / "unexpected_popup.png"))
            if w.buttons():
                w.buttons()[0].click()

    # -- tick ------------------------------------------------------------------
    def tick(self):
        try:
            self._handle_messageboxes()
            if self.sel is not None and not self.sel["done"]:
                self._tick_selection()
            if self.picker is not None and not self.picker["done"]:
                self._tick_picker()
        except Exception as exc:
            import traceback
            traceback.print_exc()
            finish(1, f"driver crashed: {exc}")

    def _tick_popup(self, box):
        st = self.popup
        now = time.time()
        if st["step"] == 0:
            st["step"], st["t"] = 1, now
        elif st["step"] == 1 and now - st["t"] > 0.7:
            cam.save(box, st["shot"])
            st["step"], st["t"] = 2, now
        elif st["step"] == 2 and now - st["t"] > 0.2:
            if box.buttons():
                box.buttons()[0].click()
            st["done"] = True

    def _tick_selection(self):
        st = self.sel
        now = time.time()
        dlg = self._find("SelectionDialog")
        if dlg is None:
            return
        if st["step"] == 0:
            st["step"], st["t"] = 1, now
        elif st["step"] == 1 and now - st["t"] > 0.8:
            cam.save(dlg, st["shot"])
            st["step"], st["t"] = 2, now
        elif st["step"] == 2 and now - st["t"] > 0.3:
            dlg.run_btn.click()
            st["done"] = True

    def _tick_picker(self):
        from PySide6.QtWidgets import QLineEdit, QListView, QPushButton
        pk = self.picker
        now = time.time()
        dlg = self._find("_ChopperDialog", "select the 922 batch folder")
        if dlg is None:
            if pk["step"] >= 5:              # closed - all done
                pk["done"] = True
            return

        def batch_index():
            view = dlg.findChild(QListView, "listView")
            if view is None:
                return None, None
            model = view.model()
            root = view.rootIndex()
            for r in range(model.rowCount(root)):
                idx = model.index(r, 0, root)
                if str(idx.data()) == f"Batch {BATCH}":
                    return view, idx
            return view, None

        if pk["step"] == 0:
            # The machine's own Qt sidebar bookmarks must not appear in a
            # picture; show only the fake root (restored before closing,
            # because the dialog persists its state).
            from PySide6.QtCore import QUrl
            pk["saved_sidebar"] = dlg.sidebarUrls()
            drive = _MOUNTED[0] if _MOUNTED else "T:"
            dlg.setSidebarUrls([QUrl.fromLocalFile(f"{drive}/{ROOT_NAME}")])
            dlg.resize(640, 430)
            geo = dlg._overlay.geometry() if dlg._overlay else None
            if geo is not None:
                dlg.move(geo.x() + (geo.width() - dlg.width()) // 2,
                         geo.y() + 110)
            pk["step"], pk["t"] = 1, now
        elif pk["step"] == 1 and now - pk["t"] > 0.4:
            view, idx = batch_index()
            if view is None or idx is None:
                return                        # model still populating - retry
            view.setCurrentIndex(idx)         # -> lock-on animation
            pk["step"], pk["t"] = 2, now
        elif pk["step"] == 2 and now - pk["t"] > 1.0:
            view, idx = batch_index()
            if idx is not None:
                dlg._on_click(idx)            # commit: TARGET CONFIRMED
            edit = dlg.findChild(QLineEdit, "fileNameEdit")
            if edit is not None:
                edit.setText(f"Batch {BATCH}")
            pk["step"], pk["t"] = 3, now
        elif pk["step"] == 3 and now - pk["t"] > 0.8:
            if pk["shot"]:
                self._composite_picker_shot(dlg, pk["shot"])
            dlg.setSidebarUrls(pk.get("saved_sidebar") or [])
            pk["step"], pk["t"] = 4, now
        elif pk["step"] == 4 and now - pk["t"] > 0.3:
            if pk["mode"] == "cancel":
                dlg.reject()                  # Cancel never fires the kill-cam
                pk["step"], pk["t"] = 5, now
                return
            for b in dlg.findChildren(QPushButton):
                if b.text().replace("&", "").strip().lower() == "execute":
                    b.click()                 # fire (kill-cam starts)
                    break
            pk["step"], pk["t"] = 5, now
        elif pk["step"] == 5 and now - pk["t"] > 0.35:
            ov = getattr(dlg, "_overlay", None)
            if ov is not None:
                ov.skip()                     # jump the kill-cam to the close
            pk["t"] = now                     # keep nudging until dlg closes

    def _composite_picker_shot(self, dlg, name: str):
        """Dialog + gunner-HUD overlay composited on the dark feed, cropped
        to the dialog plus the HUD strip above it (same look as the other
        922 chapters' picker shots)."""
        from PySide6.QtCore import QPoint
        from PySide6.QtGui import QColor, QImage, QPainter
        ov = dlg._overlay
        dpix = dlg.grab()
        if ov is None:
            dpix.save(str(cam.OUT_DIR / f"{name}.png"))
            print(f"  {name}.png (dialog only)")
            return
        ogeo = ov.geometry()
        canvas = QImage(ogeo.width(), ogeo.height(),
                        QImage.Format.Format_ARGB32_Premultiplied)
        canvas.fill(QColor(8, 10, 8))
        p = QPainter(canvas)
        tl = dlg.mapToGlobal(QPoint(0, 0))
        dx, dy = tl.x() - ogeo.x(), tl.y() - ogeo.y()
        p.drawPixmap(dx, dy, dpix)
        p.drawPixmap(0, 0, ov.grab())        # translucent HUD over everything
        p.end()
        x0 = max(0, dx - 80)
        y0 = 0
        x1 = min(ogeo.width(), dx + dpix.width() + 80)
        y1 = min(ogeo.height(), dy + dpix.height() + 56)
        crop = canvas.copy(x0, y0, x1 - x0, y1 - y0)
        cam.OUT_DIR.mkdir(parents=True, exist_ok=True)
        crop.save(str(cam.OUT_DIR / f"{name}.png"))
        print(f"  {name}.png  ({crop.width()}x{crop.height()})", flush=True)


# --- Straight-line helpers ---------------------------------------------------

def pump_until(app, window, cond, timeout_s: float, what: str):
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        app.processEvents()
        if cond():
            return
    print(f"TIMEOUT waiting for: {what}", flush=True)
    print("console tail:\n" + cam.console_tail(window, 20), flush=True)
    finish(1, f"timeout: {what}")


def run_finished(window) -> bool:
    return not window.home_page._run.session.is_running


def save_console(app, window, name: str, console_px: int = 500):
    try:
        window.home_splitter.setSizes([800 - console_px, console_px])
    except Exception:
        pass
    cam.pump(app, 700)
    cam.save(window.console, name)


def clear_console(app, window):
    try:
        window.console.clear()
    except Exception:
        pass
    cam.pump(app, 200)


# --- Scenes -------------------------------------------------------------------

def scene_repeater(app, window, settings, driver, root: Path):
    print("[922_batch_repeater]", flush=True)
    settings.set_plugin_setting("922_batch_repeater", "base_directory",
                                str(root))
    # The MASTER PARTS half needs the batch's REV C workbook (not part of this
    # fixture); off keeps the console honest instead of showing a warning.
    settings.set_plugin_setting("922_batch_repeater", "update_master_parts",
                                False)
    driver.expect_selection("922_batch_repeater_select_repeats")

    clear_console(app, window)
    cam.start_app(app, window, "922_batch_repeater")
    if not cam.wait_console_prompt(app, window, timeout_ms=45000,
                                   contains="batch number"):
        print("console tail:\n" + cam.console_tail(window, 20), flush=True)
        finish(1, "typed batch prompt never came")

    # Picture 1: the typed batch prompt, answer typed but not yet sent.
    con = window.console
    try:
        window.home_splitter.setSizes([440, 360])
    except Exception:
        pass
    con.input_field.setText(BATCH)
    cam.pump(app, 500)
    cam.save(con, "922_batch_repeater_prompt")
    con._on_input_submitted()
    cam.pump(app, 300)

    # Picture 2 (Select Repeats to Pull) is taken by the driver.
    pump_until(app, window, driver.sel_done, 60, "Select Repeats window")
    pump_until(app, window, lambda: run_finished(window), 90,
               "922 Batch Repeater run end")
    tail = cam.console_tail(window, 12)
    if "All done successfully!" not in tail:
        print("console tail:\n" + cam.console_tail(window, 30), flush=True)
        finish(1, "repeater run did not finish clean")

    # Picture 3: the finished console with the COPY SUMMARY.
    save_console(app, window, "922_batch_repeater_summary", console_px=520)


def scene_kitting(app, window, settings, driver, root: Path):
    print("[922_kitting]", flush=True)
    driver.expect_picker("922_kitting_batch_pick", mode="cancel")
    clear_console(app, window)
    cam.start_app(app, window, "922_kitting")
    pump_until(app, window, driver.picker_done, 60, "Kitting batch picker")
    pump_until(app, window, lambda: run_finished(window), 30,
               "Kitting run end (cancelled pick)")


def scene_lst(app, window, settings, driver, root: Path):
    print("[922_lst_organizer]", flush=True)
    driver.expect_picker("922_lst_organizer_batch_pick", mode="fire")
    driver.expect_popup("922 LST", "922_lst_organizer_results")
    clear_console(app, window)
    cam.start_app(app, window, "922_lst_organizer")
    pump_until(app, window, driver.picker_done, 60, "LST batch picker")
    pump_until(app, window, driver.popup_done, 90, "LST results popup")
    pump_until(app, window, lambda: run_finished(window), 30,
               "LST Organizer run end")
    print("console tail:\n" + cam.console_tail(window, 14), flush=True)


def scene_genie(app, window, settings, driver, root: Path):
    print("[922_runtime_genie]", flush=True)
    driver.expect_picker("922_runtime_genie_batch_pick", mode="fire")
    clear_console(app, window)
    cam.start_app(app, window, "922_runtime_genie")
    pump_until(app, window, driver.picker_done, 60, "Genie batch picker")
    pump_until(app, window, lambda: run_finished(window), 60,
               "Runtime Genie run end")
    save_console(app, window, "922_runtime_genie_summary", console_px=460)


SCENES = {
    "repeater": scene_repeater,
    "kitting": scene_kitting,
    "lst": scene_lst,
    "genie": scene_genie,
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scene", default="all",
                    choices=[*SCENES, "all"],
                    help="which capture scene to run (default: all)")
    args = ap.parse_args()
    wanted = list(SCENES) if args.scene == "all" else [args.scene]

    # Backstop in case even Qt timers are wedged (native loop, deadlock).
    threading.Timer(150, lambda: (print("HARD BACKSTOP (150s)", flush=True),
                                  finish(1, "threading backstop"))).start()

    # Each invocation builds ONE fixture state; running multiple scenes in a
    # row rebuilds it between scenes.
    first = True
    app = window = settings = driver = None
    root = None
    for scene in wanted:
        build_fixture(scene)
        if first:
            root = mount_drive()
            app, window, settings = cam.boot()

            # No sound, no telemetry, no real webhook - ever, in a capture.
            from techdeck.core.audio_manager import get_audio_manager
            get_audio_manager().set_enabled(False)
            import techdeck.core.usage_tracker as ut
            ut.TELEMETRY_WEBHOOK_URL = ""
            from techdeck.core import plugin_sdk as sdkmod

            def _fake_post(url, payload, log):
                log("Webhook accepted the request (HTTP 202).")
                return True
            sdkmod.post_webhook = _fake_post

            # Only this family in the kit (run_selected filters on it).
            settings.set_profile_tiles(APP_KIT)
            # Sentry Drone owned + on for the picker apps so the batch pick
            # is a drivable Qt window, never a native dialog.
            settings.unlock_item("toy_sentry_drone")
            for pid in ("922_kitting", "922_lst_organizer",
                        "922_runtime_genie"):
                settings.set_plugin_setting(pid, "sentry_drone", True)
                settings.set_plugin_setting(pid, "base_path", str(root))

            from PySide6.QtCore import QTimer
            QTimer.singleShot(120_000, lambda: (
                print("SAFETY TIMER (120s) - console tail:\n"
                      + cam.console_tail(window, 20), flush=True),
                finish(1, "global 120s safety"),
            ))
            driver = Driver(app, window)
            first = False

        SCENES[scene](app, window, settings, driver, root)

    print("Done.")
    finish(0)


if __name__ == "__main__":
    main()
