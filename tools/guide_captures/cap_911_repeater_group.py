r"""User Guide captures: 911 Batch Repeater (full treatment) + the three
simple 911 extract/prep apps (one shot each).

Run from the repo root:

    PYTHONUTF8=1 python -u tools/guide_captures/cap_911_repeater_group.py

Shots produced (docs/user_guide/images/):
  911_batch_repeater_pick.png     - phase-1 Sentry Drone batch folder pick
  911_batch_repeater_nests.png    - "Select Nests to Run" window, one nest expanded
  911_batch_repeater_summary.png  - console REPEATER SUMMARY after a real run
  911_po_pdf_extractor_prompt.png - console at its first prompt
  911_sketch_extractor_prompt.png - console at its first prompt
  911_scripting_prep_pick.png     - the award-review file pick (Sentry Drone on)

All data is FAKE (practice batches under C:\Temp\Pilot Program). The Sentry
Drone gadget is granted in the sandbox profile so folder/file picks are Qt
dialogs that can be photographed and driven; the default native Explorer
dialogs cannot be. Event-driven throughout (QTimer + generator): a modal
dialog's nested event loop would freeze straight-line code.
"""

from __future__ import annotations

import os
import shutil
import sys
import time
import traceback
from pathlib import Path

sys.path.insert(0, r"C:\Dev\TechDeck\tools")
import capture_guide_screens as cam  # noqa: E402  (sandboxes LOCALAPPDATA on import)

# ── Fake practice data (never a real batch, person, or production path) ─────
FIX_ROOT = Path(r"C:\Temp\Pilot Program")
QTDR = FIX_ROOT / "911 QTDR"
AWARD_DIR = FIX_ROOT / "Award Packages"
AWARD_FILE = AWARD_DIR / "911 SSPO AWARD REVIEW - 1000129724 SSPO Award 14.xlsx"
BATCH = "V060"

# nest -> (repeat DYPNs with their completed source, count of new DYPNs)
_REPEATS = {
    "504100": [("BK573423", "V041", "503200"), ("BK573424", "V041", "503200"),
               ("BK581101", "V041", "503200")],
    "504101": [("BK573988", "V052", "507300"), ("BK590112", "V052", "507300")],
    "504102": [("BK573423", "V041", "503200"), ("BK590113", "V052", "507300"),
               ("BK590114", "V052", "507300"), ("BK573989", "V052", "507300")],
}
_NEW_COUNTS = {"504100": 9, "504101": 6, "504102": 6}


def build_fixtures() -> None:
    from openpyxl import Workbook

    if FIX_ROOT.exists():
        shutil.rmtree(FIX_ROOT, ignore_errors=True)

    # Target batch: nest folders + their nest workbooks (NEST sheet, DYPN col).
    serial = 510
    for nest, repeats in _REPEATS.items():
        nest_dir = QTDR / BATCH / nest
        nest_dir.mkdir(parents=True)
        dypns = [d for d, _, _ in repeats]
        for _ in range(_NEW_COUNTS[nest]):
            dypns.append(f"BK600{serial}")
            serial += 1
        wb = Workbook()
        ws = wb.active
        ws.title = "NEST"
        for c, h in enumerate(["ITEM", "DYPN", "QTY", "MATL"], 1):
            ws.cell(1, c, h)
        for r, dypn in enumerate(dypns, 2):
            ws.cell(r, 1, r - 1)
            ws.cell(r, 2, dypn)
            ws.cell(r, 3, 1)
            ws.cell(r, 4, "A36 PLATE")
        wb.save(nest_dir / f"911 BATCH {BATCH} {nest}.xlsx")

    # One nest already has REPEAT files, so the window shows its flag.
    staged = QTDR / BATCH / "504102" / "REPEAT" / "BK590113"
    staged.mkdir(parents=True)
    (staged / "BK590113.pdf").write_bytes(b"fake pdf")

    # Completed source nests: CAD-AND-SHOP-PRINTS\{DYPN}\ with the real trio.
    sources = {(d, b, n) for reps in _REPEATS.values() for d, b, n in reps}
    for dypn, batch, nest in sorted(sources):
        part_dir = QTDR / batch / nest / "CAD-AND-SHOP-PRINTS" / dypn
        part_dir.mkdir(parents=True, exist_ok=True)
        for ext in (".SLDPRT", ".SLDDRW", ".pdf"):
            (part_dir / f"{dypn}{ext}").write_bytes(b"fake " + ext.encode())

    # The 911 MASTER PARTS LIST in the REPEATER folder.
    mpl_dir = (QTDR / "04 - Notes - Protocols - Tutorials" / "TECH SERVICES"
               / "REPEATER")
    mpl_dir.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "MASTER PARTS"
    for c, h in enumerate(
            ["DYPN", "BATCH", "NEST", "STATUS", "SOURCE FOLDER", "FILES"], 1):
        ws.cell(1, c, h)
    for r, (dypn, batch, nest) in enumerate(sorted(sources), 2):
        ws.cell(r, 1, dypn)
        ws.cell(r, 2, batch)
        ws.cell(r, 3, nest)
        ws.cell(r, 4, "SHIPPED")
        ws.cell(r, 5, f"{batch}/{nest}/CAD-AND-SHOP-PRINTS/{dypn}")
        ws.cell(r, 6, "pdf, sldprt, slddrw")
    wb.save(mpl_dir / "911 MASTER PARTS LIST.xlsx")

    # A fake award review for the 911 Scripting Prep file pick.
    AWARD_DIR.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Working Forecast Input"
    for c, h in enumerate(["Order", "Source Material", "Nest Pkg Nbr"], 1):
        ws.cell(1, c, h)
    ws.cell(2, 1, "BK612001")
    ws.cell(2, 2, "PL0500A36")
    ws.cell(2, 3, "509400")
    wb.save(AWARD_FILE)
    print(f"Fixtures staged under {FIX_ROOT}", flush=True)


# ── Event-driven state machine ───────────────────────────────────────────────
# The script body is a generator; each `yield (condition, timeout_s, label)`
# suspends it until the QTimer tick sees the condition come true (probe()'s
# pattern - a modal dialog's nested event loop still runs QTimers, so this
# keeps working while straight-line pumps would freeze).

_app = None
_window = None
_settings = None
_gen = None
_cond = None
_deadline = 0.0
_label = ""


def _die(msg: str) -> None:
    print(f"CAPTURE FAILED: {msg}", flush=True)
    try:
        print("-- console tail:", flush=True)
        print(cam.console_tail(_window, 20), flush=True)
    except Exception:
        pass
    sys.stdout.flush()
    os._exit(1)


def _advance(value=None) -> None:
    global _cond, _deadline, _label
    try:
        step = _gen.send(value)
    except StopIteration:
        print("All captures done.", flush=True)
        sys.stdout.flush()
        os._exit(0)
    except Exception:
        traceback.print_exc()
        _die("script step raised")
    _cond, timeout_s, _label = step
    _deadline = time.time() + timeout_s


def _tick() -> None:
    global _cond
    if _cond is None:
        return
    try:
        result = _cond()
    except Exception:
        traceback.print_exc()
        _die(f"condition raised in step: {_label}")
        return
    if result:
        _cond = None
        _advance(result)
    elif time.time() > _deadline:
        _die(f"timed out waiting for: {_label}")


# ── Small helpers used by the script generator ───────────────────────────────

def _start(plugin_id: str) -> None:
    # cam.start_app minus the trailing pump: a pump here can get trapped in a
    # dialog's nested event loop before the generator reaches its next yield.
    _window.home_page.selected_tiles = {plugin_id}
    _window.home_page._run.run_selected_plugins()


def _find_window(type_name: str, title_contains: str = ""):
    from PySide6.QtWidgets import QApplication

    def cond():
        for w in QApplication.topLevelWidgets():
            if w is _window or not w.isVisible():
                continue
            if type(w).__name__ != type_name:
                continue
            if title_contains.lower() in (w.windowTitle() or "").lower():
                return w
        return None
    return cond


def _ticks(n: int):
    """Condition that is true after n further timer ticks (layout settling)."""
    box = {"left": n}

    def cond():
        box["left"] -= 1
        return box["left"] <= 0
    return cond


def _not_running():
    return lambda: not _window.home_page._run.session.is_running


def _prompt_contains(text: str):
    con = _window.console

    def cond():
        return (con.waiting_for_input
                and text.lower() in (getattr(con, "input_prompt", "") or "").lower())
    return cond


def _cancel_run() -> None:
    if _window.home_page._run.session.is_running:
        _window.home_page._run.run_selected_plugins()   # Run doubles as Cancel


def _dress_picker(dlg) -> None:
    """Make a chopper picker photograph well: hide the squeezed sidebar
    (its clipped labels read as a rendering bug) and give the list room."""
    from PySide6.QtWidgets import QWidget

    sidebar = dlg.findChild(QWidget, "sidebar")
    if sidebar is not None:
        sidebar.hide()
    dlg.resize(720, 470)


def _select_row(dlg, name: str) -> bool:
    """Highlight the row called `name` in the picker's file list and fill the
    name box, the way a real click on the row does."""
    from PySide6.QtWidgets import QLineEdit, QListView

    view = dlg.findChild(QListView, "listView")
    if view is None:
        return False
    model = view.model()
    root = view.rootIndex()
    for r in range(model.rowCount(root)):
        idx = model.index(r, 0, root)
        if str(idx.data()) == name:
            view.setCurrentIndex(idx)
            view.scrollTo(idx)
            edit = dlg.findChild(QLineEdit, "fileNameEdit")
            if edit is not None:
                edit.setText(name)
            return True
    return False


def _clear_console() -> None:
    _window.console.clear_btn.click()


def _script():
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QPushButton
    from techdeck.core import plugin_sdk as sdk

    # ═══ 1. Batch Repeater, classic flow: nest window + summary ═════════════
    # The batch-folder pick is a native Explorer dialog with the drone off, so
    # it cannot be photographed or driven. Skip just the pick by having
    # sdk.request_directory hand back the practice batch folder; everything
    # after it (scan, window, copies, summary) is the real app running.
    real_request_directory = sdk.request_directory
    sdk.request_directory = lambda params, title="", start_dir="", style=None: \
        str(QTDR / BATCH)
    _start("911_batch_repeater")

    dlg = yield (_find_window("GroupedToggleDialog", "911 Batch Repeater"),
                 30, "nest-selection window")
    dlg._parents["504100"].setCheckState(0, Qt.CheckState.Checked)
    dlg._parents["504101"].setCheckState(0, Qt.CheckState.Checked)
    dlg._parents["504100"].setExpanded(True)
    yield (_ticks(3), 10, "nest window settle")
    cam.save(dlg, "911_batch_repeater_nests")

    run_btn = next(b for b in dlg.findChildren(QPushButton)
                   if b.text() == "Run Repeater")
    run_btn.click()
    yield (_not_running(), 40, "repeater run to finish")
    sdk.request_directory = real_request_directory

    # Tall console so the whole REPEATER SUMMARY block is on screen.
    _window.home_splitter.setSizes([220, 580])
    yield (_ticks(3), 10, "console relayout")
    cam.save(_window.console, "911_batch_repeater_summary")
    _window.home_splitter.setSizes([520, 280])
    yield (_ticks(2), 10, "console restore")

    # ═══ 2. PO PDF Extractor: first console prompt ═══════════════════════════
    _clear_console()
    _start("911_po_pdf_extractor")
    yield (_prompt_contains("PO packet PDFs"), 40, "PO extractor prompt")
    yield (_ticks(2), 10, "PO prompt settle")
    cam.save(_window.console, "911_po_pdf_extractor_prompt")
    _cancel_run()
    yield (_not_running(), 30, "PO extractor cancel")

    # ═══ 3. Sketch Extractor: first console prompt ═══════════════════════════
    _clear_console()
    _start("911_sketch_extractor")
    yield (_prompt_contains("containing PDFs"), 40, "sketch extractor prompt")
    yield (_ticks(2), 10, "sketch prompt settle")
    cam.save(_window.console, "911_sketch_extractor_prompt")
    _cancel_run()
    yield (_not_running(), 30, "sketch extractor cancel")

    # ═══ 4. Scripting Prep: the award-review pick (Sentry Drone = Qt dialog) ═
    _clear_console()
    _settings.set_plugin_setting("911_scripting_prep", "sentry_drone", True)
    _start("911_scripting_prep")
    dlg = yield (_find_window("_ChopperDialog", "SSPO AWARD REVIEW"),
                 40, "scripting prep file pick")
    _dress_picker(dlg)
    dlg.setDirectory(str(AWARD_DIR))
    yield (_ticks(3), 10, "file model populate")
    yield (lambda d=dlg: _select_row(d, AWARD_FILE.name), 10, "award row select")
    yield (_ticks(2), 10, "file pick settle")
    cam.save(dlg, "911_scripting_prep_pick")
    dlg.reject()
    yield (_not_running(), 30, "scripting prep cancel")

    # ═══ 5. Batch Repeater: the batch folder pick (Sentry Drone = Qt dialog) ═
    _settings.set_plugin_setting("911_batch_repeater", "sentry_drone", True)
    _start("911_batch_repeater")
    dlg = yield (_find_window("_ChopperDialog", "Select the 911 batch folder"),
                 40, "repeater batch pick")
    _dress_picker(dlg)
    yield (_ticks(3), 10, "folder model populate")
    yield (lambda d=dlg: _select_row(d, BATCH), 10, "batch row select")
    yield (_ticks(2), 10, "batch pick settle")
    cam.save(dlg, "911_batch_repeater_pick")
    dlg.reject()
    yield (_not_running(), 30, "repeater pick cancel")


def main() -> None:
    global _app, _window, _settings, _gen

    build_fixtures()
    _app, _window, _settings = cam.boot()

    # No gunship ambience / chimes out of the real speakers during capture.
    from techdeck.core.audio_manager import get_audio_manager
    get_audio_manager().set_enabled(False)
    _settings.set_audio_settings(False, 0)

    # The queue only runs tiles that are in the Home kit.
    _settings.set_profile_tiles([
        "911_setup", "911_batch_repeater", "911_po_pdf_extractor",
        "911_scripting_prep", "911_sketch_extractor", "911_remove_ticket",
    ])

    # Point the repeater at the practice tree; own the Sentry Drone gadget
    # (each shot switches it on per app only when that shot needs it).
    _settings.set_plugin_setting("911_batch_repeater", "qtdr_base_path", str(QTDR))
    _settings.unlock_item("toy_sentry_drone")

    from PySide6.QtCore import QTimer

    # Global safety net: a native dialog (or anything else) hanging the flow
    # must never hang the terminal.
    safety = QTimer()
    safety.setSingleShot(True)
    safety.timeout.connect(lambda: _die("global 90s safety timer"))
    safety.start(90_000)

    ticker = QTimer()
    ticker.timeout.connect(_tick)
    ticker.start(350)

    _gen = _script()
    _advance()          # run to the first yield
    _app.exec()
    os._exit(1)         # unreachable: the generator exits the process


if __name__ == "__main__":
    main()
