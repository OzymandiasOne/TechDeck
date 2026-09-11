"""User Guide captures: 922 Setup, 922 Pallet Stamper, 922 Difficulty Stamper,
922 FormingFinder.

Run from the repo root:

    PYTHONUTF8=1 python -u tools/guide_captures/cap_922_setup_family.py
    ... --apps 922_setup                      # retake one app only

Everything runs against the camera library's fresh sandbox profile with a FAKE
922 batch tree (Batch 573, invented order numbers) mounted on a subst'd T:
drive so no real path, batch, or user name can appear in a picture.

The Sentry Drone gadget is unlocked and switched on for every app so the
batch-folder pick opens as a Qt window (the chopper picker) instead of the
native Explorer dialog -- a native dialog can neither be photographed nor
driven and would hang the run.

All modal dialogs (the stage toggle window, the drone picker, any stray
QMessageBox) are driven from a QTimer state machine: a modal exec() nests the
event loop, so straight-line code freezes until the dialog closes, but timers
keep firing inside the nested loop. A global safety timer hard-exits after
120 s no matter what.
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, r"C:\Dev\TechDeck\tools")
import capture_guide_screens as cam  # noqa: E402  (sets the sandbox FIRST)

BATCH = "573"
ROOT_NAME = "922 QTDR Production Packages"
DRIVE = "T:"
FAKE_PARENT = cam._SANDBOX / "fake_shop"

# (order, ppn, pallet column letter is decided in the organizer below)
ORDERS = [
    ("BK573423", "R7651664-H23"),
    ("BK573424", "R7651665-H23"),
    ("BK573425", "R7651666-H23"),
    ("X6401069", "R7651667-H23"),
]

APP_IDS = ["922_setup", "922_pallet_stamper", "922_difficulty_stamper",
           "922_formingfinder"]


def _subst_off():
    subprocess.run(f"subst {DRIVE} /D", shell=True, capture_output=True)


def finish(code: int, why: str = ""):
    if why:
        print(f"-- exit({code}): {why}", flush=True)
    _subst_off()
    sys.stdout.flush()
    os._exit(code)


# ─── Fixtures ────────────────────────────────────────────────────────────────

def _make_work_packet(path: Path, order: str):
    """A one-page PDF that reads as a WORK PACKET (no title-block markers)."""
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 80), "WORK PACKET", fontsize=26, fontname="hebo")
    page.insert_text((72, 120), f"Order: {order}", fontsize=14)
    page.insert_text((72, 142), f"Batch: {BATCH}", fontsize=14)
    page.insert_text((72, 186), "Work Instruction", fontsize=12, fontname="hebo")
    y = 210
    for line in ("Lead Trade: Fabrication",
                 "1. Verify material against the PO line item.",
                 "2. Cut per program. Deburr all edges.",
                 "3. Route to inspection with this packet."):
        page.insert_text((72, y), line, fontsize=11)
        y += 20
    doc.save(str(path))
    doc.close()


def _make_drawing(path: Path, part: str, difficult: bool):
    """A one-page PDF that reads as a part DRAWING (two title-block markers),
    optionally carrying the DriveWorks-style DIFFICULT label as its own span."""
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=792, height=612)  # landscape print
    page.insert_text((60, 70), f"PART: {part}", fontsize=16, fontname="hebo")
    page.insert_text((60, 95), f"BATCH {BATCH}", fontsize=10)
    page.draw_rect(fitz.Rect(60, 130, 540, 480))
    if difficult:
        page.insert_text((600, 90), "DIFFICULT", fontsize=18,
                         fontname="hebo", fill=(0, 0, 0.85))
    page.insert_text((60, 560), "UNLESS OTHERWISE SPECIFIED", fontsize=8)
    page.insert_text((300, 560), "DO NOT SCALE DRAWING", fontsize=8)
    page.insert_text((520, 560), "ENG APPR", fontsize=8)
    doc.save(str(path))
    doc.close()


def build_fixtures():
    from openpyxl import Workbook

    root = FAKE_PARENT / ROOT_NAME
    batch = root / f"Batch {BATCH}"
    doc_dir = batch / f"Batch {BATCH} - Documentation"
    wp_dir = batch / "Work Packets"
    for d in (doc_dir, wp_dir):
        d.mkdir(parents=True, exist_ok=True)

    # REV C PO workbook: ORDER/PPN pairs on the PO sheet (headers on row 3).
    wb = Workbook()
    ws = wb.active
    ws.title = "PO"
    ws["A1"] = f"PO H{BATCH}  QF-QU-09 REV C"
    ws["A3"], ws["B3"], ws["C3"] = "ITEM", "ORDER", "PPN"
    ws["D3"], ws["E3"], ws["F3"] = "DESCRIPTION", "QTY", "NOTES"
    descs = ["PLATE, FORMED", "FLAT BAR", "PLATE", "TUBE ASSEMBLY"]
    for i, (order, ppn) in enumerate(ORDERS):
        r = 4 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, order)
        ws.cell(r, 3, ppn)
        ws.cell(r, 4, descs[i])
        ws.cell(r, 5, i + 1)
    wb.save(doc_dir / f"PO H{BATCH} QF-QU-09 REV C.xlsx")

    # Pallet & Rod Organizer: Pallet Organizer sheet, headers on Excel row 4,
    # pallet columns B / E / H (the stamper reads usecols B,E,H, header=3).
    wb = Workbook()
    ws = wb.active
    ws.title = "Pallet Organizer"
    ws["B2"] = f"BATCH {BATCH} PALLET ORGANIZER"
    ws["B4"], ws["E4"], ws["H4"] = "PALLET 1", "PALLET 2", "PALLET 3"
    ws["B5"] = "BK573424"
    ws["E5"] = "BK573423"
    ws["E6"] = "BK573425"
    ws["H5"] = "X6401069"
    wb.save(doc_dir / f"PO H{BATCH} Pallet & Rod Organizer.xlsx")

    # Work-packet PDFs waiting in the Work Packets drop folder (922 Setup's
    # folder stage files them into the order folders it builds).
    names = {"BK573423": "BK573423.pdf", "BK573424": "BK573424.pdf",
             "BK573425": "BK573425.pdf", "X6401069": "X6401069 NOFORN.pdf"}
    for order, _ppn in ORDERS:
        _make_work_packet(wp_dir / names[order], order)

    # Existing non-empty Forming folder so FormingFinder asks its
    # overwrite question.
    forming = doc_dir / f"Forming {BATCH}"
    forming.mkdir(exist_ok=True)
    import fitz
    doc = fitz.open()
    doc.new_page()
    doc.save(str(forming / f"Forming {BATCH}.pdf"))
    doc.close()

    print(f"Fixtures: {batch}")


def add_cad_prints():
    """After 922 Setup built the order folders: drop part drawings under each
    order's CAD-AND-SHOP-PRINTS folder. BK573423 gets one DIFFICULT part."""
    root = Path(f"{DRIVE}\\{ROOT_NAME}")
    batch = root / f"Batch {BATCH}"
    for order, ppn in ORDERS:
        folder = batch / f"{order}-{ppn}"
        if not folder.is_dir():
            print(f"!! missing order folder {folder} - folder stage failed?")
            finish(1, "order folders absent after 922 Setup")
        cad = folder / "CAD-AND-SHOP-PRINTS"
        cad.mkdir(exist_ok=True)
        _make_drawing(cad / f"{ppn} PLT 1.pdf", f"{ppn} PLT 1",
                      difficult=(order == "BK573423"))
        _make_drawing(cad / f"{ppn} PLT 2.pdf", f"{ppn} PLT 2", difficult=False)
    # The now-empty Work Packets drop folder would read as a fifth "order"
    # to the stampers and put a warning in every picture - drop it.
    wp_dir = batch / "Work Packets"
    if wp_dir.is_dir() and not any(wp_dir.iterdir()):
        wp_dir.rmdir()
    print("CAD-AND-SHOP-PRINTS drawings staged.")


# ─── Qt driver (runs inside nested modal event loops via QTimer) ────────────

class Driver:
    """One QTimer tick loop that answers whatever modal window is up.

    ``expect_stage(...)`` / ``expect_picker(...)`` arm a handler; the tick
    walks that handler's steps with real-time delays. Stray QMessageBoxes are
    always dismissed so no unexpected popup can hang the run.
    """

    def __init__(self, app, window):
        self.app = app
        self.window = window
        self.stage = None    # dict while driving the Select Stages window
        self.picker = None   # dict while driving the drone folder picker
        from PySide6.QtCore import QTimer
        self.timer = QTimer()
        self.timer.timeout.connect(self.tick)
        self.timer.start(150)

    # -- arming ------------------------------------------------------------
    def expect_stage(self, shot: str | None, keep: set):
        """Drive the 922 Setup Select Stages window: photograph it (default
        state, option groups unfolded), keep only ``keep`` stages checked,
        click Run Selected."""
        self.stage = {"step": 0, "t": 0.0, "shot": shot, "keep": keep,
                      "done": False}

    def expect_picker(self, shots: list):
        """Drive the drone batch picker: lock Batch 573, photograph the
        composited HUD+dialog under each name in ``shots``, fire, skip the
        kill-cam."""
        self.picker = {"step": 0, "t": 0.0, "shots": list(shots),
                       "dlg": None, "done": False}

    def stage_done(self):
        return self.stage is not None and self.stage["done"]

    def picker_done(self):
        return self.picker is not None and self.picker["done"]

    # -- helpers -----------------------------------------------------------
    def _find(self, cls_name: str, title_sub: str = ""):
        from PySide6.QtWidgets import QApplication
        for w in QApplication.topLevelWidgets():
            if (w.isVisible() and type(w).__name__ == cls_name
                    and title_sub.lower() in (w.windowTitle() or "").lower()):
                return w
        return None

    def _dismiss_messageboxes(self):
        from PySide6.QtWidgets import QApplication, QMessageBox
        for w in QApplication.topLevelWidgets():
            if isinstance(w, QMessageBox) and w.isVisible():
                print(f"!! unexpected popup: {w.windowTitle()!r}: "
                      f"{w.text()[:300]}", flush=True)
                w.grab().save(str(FAKE_PARENT / "unexpected_popup.png"))
                btns = w.buttons()
                if btns:
                    btns[0].click()

    # -- tick --------------------------------------------------------------
    def tick(self):
        try:
            self._dismiss_messageboxes()
            if self.stage is not None and not self.stage["done"]:
                self._tick_stage()
            if self.picker is not None and not self.picker["done"]:
                self._tick_picker()
        except Exception as exc:  # never let the driver die silently
            import traceback
            traceback.print_exc()
            finish(1, f"driver crashed: {exc}")

    def _tick_stage(self):
        from PySide6.QtCore import Qt
        st = self.stage
        now = time.time()
        dlg = self._find("GroupedToggleDialog")
        if dlg is None:
            return
        if st["step"] == 0:
            # Unfold the groups that have option rows so the picture shows
            # the stage toggles AND their options (a state the user reaches
            # by clicking the stage names).
            for gkey, parent in dlg._parents.items():
                if parent.childCount():
                    parent.setExpanded(True)
            st["step"], st["t"] = 1, now
        elif st["step"] == 1 and now - st["t"] > 0.7:
            if st["shot"]:
                cam.save(dlg, st["shot"])
            st["step"], st["t"] = 2, now
        elif st["step"] == 2 and now - st["t"] > 0.2:
            for gkey, parent in dlg._parents.items():
                want = Qt.CheckState.Checked if gkey in st["keep"] \
                    else Qt.CheckState.Unchecked
                parent.setCheckState(0, want)
            st["step"], st["t"] = 3, now
        elif st["step"] == 3 and now - st["t"] > 0.4:
            dlg.run_btn.click()
            st["done"] = True

    def _tick_picker(self):
        from PySide6.QtCore import QPoint
        from PySide6.QtWidgets import QListView, QPushButton
        pk = self.picker
        now = time.time()
        dlg = self._find("_ChopperDialog", "select the 922 batch folder")
        if dlg is None:
            if pk["step"] >= 5:            # fired and closed - all done
                pk["done"] = True
            return
        pk["dlg"] = dlg

        def batch_index():
            view = dlg.findChild(QListView, "listView")
            if view is None:
                return None, None
            model = view.model()
            root = view.rootIndex()
            for r in range(model.rowCount(root)):
                idx = model.index(r, 0, root)
                if str(idx.data()) == f"Batch {BATCH}":
                    return view, idx
            return view, None

        if pk["step"] == 0:
            # The machine's own Qt sidebar bookmarks must not appear in a
            # picture; remember them (restored before accept, because accept
            # persists the dialog state) and show only the fake root.
            from PySide6.QtCore import QUrl
            pk["saved_sidebar"] = dlg.sidebarUrls()
            dlg.setSidebarUrls([QUrl.fromLocalFile(
                f"{DRIVE}/{ROOT_NAME}".replace("\\", "/"))])
            dlg.resize(640, 430)
            geo = dlg._overlay.geometry() if dlg._overlay else None
            if geo is not None:
                dlg.move(geo.x() + (geo.width() - dlg.width()) // 2,
                         geo.y() + 110)
            pk["step"], pk["t"] = 1, now
        elif pk["step"] == 1 and now - pk["t"] > 0.4:
            view, idx = batch_index()
            if view is None or idx is None:
                return                      # model still populating - retry
            view.setCurrentIndex(idx)       # -> lock-on animation
            pk["step"], pk["t"] = 2, now
        elif pk["step"] == 2 and now - pk["t"] > 1.0:
            view, idx = batch_index()
            if idx is not None:
                dlg._on_click(idx)          # commit: TARGET CONFIRMED
            # A real click also fills the name field with the folder name.
            from PySide6.QtWidgets import QLineEdit
            edit = dlg.findChild(QLineEdit, "fileNameEdit")
            if edit is not None:
                edit.setText(f"Batch {BATCH}")
            pk["step"], pk["t"] = 3, now
        elif pk["step"] == 3 and now - pk["t"] > 0.8:
            if pk["shots"]:
                self._composite_picker_shot(dlg, pk["shots"])
            dlg.setSidebarUrls(pk.get("saved_sidebar") or [])
            pk["step"], pk["t"] = 4, now
        elif pk["step"] == 4 and now - pk["t"] > 0.3:
            for b in dlg.findChildren(QPushButton):
                if b.text().replace("&", "").strip().lower() == "execute":
                    b.click()               # fire (kill-cam starts)
                    break
            pk["step"], pk["t"] = 5, now
        elif pk["step"] == 5 and now - pk["t"] > 0.35:
            ov = getattr(dlg, "_overlay", None)
            if ov is not None:
                ov.skip()                   # jump the kill-cam to the close
            pk["t"] = now                   # keep nudging until dlg closes

    def _composite_picker_shot(self, dlg, names: list):
        """Dialog + gunner-HUD overlay composited on the dark feed, cropped
        to the dialog plus the HUD strip above it."""
        from PySide6.QtCore import QPoint
        from PySide6.QtGui import QColor, QImage, QPainter
        ov = dlg._overlay
        dpix = dlg.grab()
        if ov is None:
            for n in names:
                dpix.save(str(cam.OUT_DIR / f"{n}.png"))
                print(f"  {n}.png (dialog only)")
            return
        ogeo = ov.geometry()
        canvas = QImage(ogeo.width(), ogeo.height(),
                        QImage.Format.Format_ARGB32_Premultiplied)
        canvas.fill(QColor(8, 10, 8))
        p = QPainter(canvas)
        tl = dlg.mapToGlobal(QPoint(0, 0))
        dx, dy = tl.x() - ogeo.x(), tl.y() - ogeo.y()
        p.drawPixmap(dx, dy, dpix)
        p.drawPixmap(0, 0, ov.grab())      # translucent HUD over everything
        p.end()
        x0 = max(0, dx - 80)
        y0 = 0
        x1 = min(ogeo.width(), dx + dpix.width() + 80)
        y1 = min(ogeo.height(), dy + dpix.height() + 56)
        crop = canvas.copy(x0, y0, x1 - x0, y1 - y0)
        cam.OUT_DIR.mkdir(parents=True, exist_ok=True)
        for n in names:
            crop.save(str(cam.OUT_DIR / f"{n}.png"))
            print(f"  {n}.png  ({crop.width()}x{crop.height()})", flush=True)


# ─── Straight-line helpers ───────────────────────────────────────────────────

def pump_until(app, window, cond, timeout_s: float, what: str):
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        app.processEvents()
        if cond():
            return
    print(f"TIMEOUT waiting for: {what}", flush=True)
    print("console tail:\n" + cam.console_tail(window, 20), flush=True)
    finish(1, f"timeout: {what}")


def run_finished(window) -> bool:
    return not window.home_page._run.session.is_running


def save_console(app, window, name: str):
    """Give the console some height, let the log drain, grab it."""
    try:
        window.home_splitter.setSizes([300, 500])
    except Exception:
        pass
    cam.pump(app, 700)
    cam.save(window.console, name)


def clear_console(app, window):
    """Fresh console per app so a shot never shows the previous app's run."""
    try:
        window.console.clear()
    except Exception:
        pass
    cam.pump(app, 200)


# ─── Per-app capture flows ───────────────────────────────────────────────────

def cap_922_setup(app, window, driver):
    print("[922_setup]", flush=True)
    driver.expect_stage(shot="922_setup_stages", keep={"folder_setup"})
    driver.expect_picker(shots=["922_setup_batch_pick", "run_drone_picker"])
    clear_console(app, window)
    cam.start_app(app, window, "922_setup")
    pump_until(app, window, driver.stage_done, 30, "922 Setup stage window")
    pump_until(app, window, driver.picker_done, 30, "922 Setup batch picker")
    pump_until(app, window, lambda: run_finished(window), 60,
               "922 Setup run end")
    save_console(app, window, "922_setup_run_console")


def cap_pallet_stamper(app, window, driver):
    print("[922_pallet_stamper]", flush=True)
    driver.stage = None
    driver.expect_picker(shots=["922_pallet_stamper_batch_pick"])
    clear_console(app, window)
    cam.start_app(app, window, "922_pallet_stamper")
    pump_until(app, window, driver.picker_done, 30, "Pallet Stamper picker")
    pump_until(app, window, lambda: run_finished(window), 60,
               "Pallet Stamper run end")
    save_console(app, window, "922_pallet_stamper_summary")


def cap_difficulty_stamper(app, window, driver):
    print("[922_difficulty_stamper]", flush=True)
    driver.expect_picker(shots=[])          # drive it, no photo (same look)
    clear_console(app, window)
    cam.start_app(app, window, "922_difficulty_stamper")
    pump_until(app, window, driver.picker_done, 30, "Difficulty picker")
    pump_until(app, window, lambda: run_finished(window), 60,
               "Difficulty Stamper run end")
    save_console(app, window, "922_difficulty_stamper_summary")


def cap_formingfinder(app, window, driver):
    print("[922_formingfinder]", flush=True)
    driver.expect_picker(shots=[])          # drive it, no photo (same look)
    clear_console(app, window)
    cam.start_app(app, window, "922_formingfinder")
    pump_until(app, window, driver.picker_done, 30, "FormingFinder picker")
    ok = cam.wait_console_prompt(app, window, timeout_ms=20000,
                                 contains="Overwrite existing")
    if not ok:
        print("console tail:\n" + cam.console_tail(window, 20), flush=True)
        finish(1, "FormingFinder overwrite prompt never came")
    save_console(app, window, "922_formingfinder_overwrite_prompt")
    cam.answer_console(app, window, "N")    # abort: forming preserved
    pump_until(app, window, lambda: run_finished(window), 30,
               "FormingFinder run end")


FLOWS = {
    "922_setup": cap_922_setup,
    "922_pallet_stamper": cap_pallet_stamper,
    "922_difficulty_stamper": cap_difficulty_stamper,
    "922_formingfinder": cap_formingfinder,
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apps", default=",".join(APP_IDS),
                    help="comma list of app ids to capture")
    args = ap.parse_args()
    wanted = [a.strip() for a in args.apps.split(",") if a.strip()]
    for a in wanted:
        if a not in FLOWS:
            sys.exit(f"unknown app id: {a}")

    build_fixtures()
    _subst_off()
    r = subprocess.run(f'subst {DRIVE} "{FAKE_PARENT}"', shell=True,
                       capture_output=True, text=True)
    root = Path(f"{DRIVE}\\{ROOT_NAME}")
    if not root.is_dir():
        sys.exit(f"subst failed: {r.stdout} {r.stderr}")
    print(f"Fake root mounted: {root}")

    app, window, settings = cam.boot()

    # Only the four target apps in the kit (run_selected filters on it).
    settings.set_profile_tiles(APP_IDS)
    # Sentry Drone: owned + on for every app => the batch pick is a Qt window.
    settings.unlock_item("toy_sentry_drone")
    for pid in APP_IDS:
        settings.set_plugin_setting(pid, "sentry_drone", True)
        settings.set_plugin_setting(pid, "base_path", str(root))
    # Belt and braces: 922 Setup never posts to Teams even if a card stage
    # were somehow selected.
    settings.set_plugin_setting("922_setup", "dry_run", True)

    from PySide6.QtCore import QTimer
    QTimer.singleShot(120_000, lambda: (
        print("SAFETY TIMER (120s) - console tail:\n"
              + cam.console_tail(window, 20), flush=True),
        finish(1, "global 120s safety"),
    ))

    driver = Driver(app, window)

    if "922_setup" in wanted:
        cap_922_setup(app, window, driver)
        add_cad_prints()
    else:
        # A standalone retake of a later app skips the 922 Setup run that
        # normally builds the order folders - stage them directly instead.
        batch = root / f"Batch {BATCH}"
        wp_dir = batch / "Work Packets"
        for order, ppn in ORDERS:
            folder = batch / f"{order}-{ppn}"
            folder.mkdir(exist_ok=True)
            for pdf in list(wp_dir.glob("*.pdf")):
                if order in pdf.stem:
                    pdf.rename(folder / pdf.name)
        add_cad_prints()
    if "922_pallet_stamper" in wanted:
        cap_pallet_stamper(app, window, driver)
    if "922_difficulty_stamper" in wanted:
        cap_difficulty_stamper(app, window, driver)
    if "922_formingfinder" in wanted:
        cap_formingfinder(app, window, driver)

    print("Done.")
    finish(0)


if __name__ == "__main__":
    main()
