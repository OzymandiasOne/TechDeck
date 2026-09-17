"""911 SSPO Invoicing Prep v2.3.0 - invoice number + ship date on the supplement,
and the supplement-to-PDF step (invoicing's ask 2026-09-04, answers 2026-09-16).

Pins:
  - the forecast lookup carries PS/Inv + Ship Date, found by header NAME on both
    sheets ('911 Forecast' first, 'Complete 911 QTDR' fallback);
  - the Invoice Supplement gets G2 = PS/Inv and G3 = Ship Date; blank stays blank;
  - the PDF is printed ONLY for a nest with an invoice number, named
    'ASA Invoice No. {inv} Supplement.pdf' beside its back-up workbook;
  - a nest without PS/Inv lands in missing_invoice (no PDF), one with no forecast
    row at all stays in missing_po;
  - an exporter that can't run leaves pdf_error set and still writes every workbook.
"""
import importlib.util
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

PLUGIN = (Path(__file__).resolve().parents[2]
          / "plugins" / "911_sspo_invoicing_prep" / "run.py")


@pytest.fixture(scope="module")
def mod():
    spec = importlib.util.spec_from_file_location("sp_911_sspo_invoicing_prep", PLUGIN)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


PRICING_HEADERS = ["Batch", "Work Order", "DYPN", "Material", "DYPN QTY",
                   "Nest Pkg Nbr", "TOTAL PRICE PER WO", "Firm VPD", "Machine",
                   "Scheduling Group", "Program", "SCOPE OF WORK", "SubGroup",
                   "Division", "Notes", "Material Status"]


def _pricing_workbook(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(PRICING_HEADERS)
    for batch, nest, wo, total in rows:
        ws.append([batch, wo, f"H{wo}", "HSS", 2, nest, total,
                   datetime(2026, 9, 4), "V807", "Open", "911", "CUT", "A",
                   "SOPO", "", "OK"])
    wb.save(path)


def _forecast_workbook(path: Path, active_rows, complete_rows, with_inv=True):
    """Both 911 sheets with the REAL column positions that matter: PO/Line/Batch/
    Nest up front, Ship Date at AT (col 46) and PS/Inv at BD (col 56) - the
    lookup must find them by name, so the test also proves it isn't reading a
    fixed letter."""
    wb = Workbook()
    ws = wb.active
    ws.title = "911 Forecast"
    wb.create_sheet("Complete 911 QTDR")

    def fill(sheet, batch_header, rows):
        hdr = [""] * 56
        hdr[0], hdr[1], hdr[2], hdr[6] = "PO", "Line", batch_header, "Nest"
        hdr[45] = "Ship Date"
        if with_inv:
            hdr[55] = "PS/Inv"
        sheet.append(hdr)
        for po, line, batch, nest, ship, inv in rows:
            r = [None] * 56
            r[0], r[1], r[2], r[6] = po, line, batch, nest
            r[45], r[55] = ship, inv
            sheet.append(r)

    fill(wb["911 Forecast"], "Batch /DR", active_rows)
    fill(wb["Complete 911 QTDR"], "Batch/Order", complete_rows)
    wb.save(path)


class FakeExporter:
    """Stands in for Excel: records every export and writes an empty PDF."""

    def __init__(self, fail=False, error=None):
        self.calls = []
        self.fail = fail
        self.error = error
        self.closed = False

    def export(self, xlsx_path, pdf_path):
        self.calls.append((Path(xlsx_path), Path(pdf_path)))
        if self.fail:
            return False
        Path(pdf_path).write_bytes(b"%PDF-1.4 fake")
        return True

    def close(self):
        self.closed = True


@pytest.fixture
def run_split(mod, tmp_path):
    def _run(active_rows, complete_rows, pricing_rows, exporter=None, with_inv=True):
        fdir = tmp_path / "forecast"
        fdir.mkdir()
        _forecast_workbook(fdir / "Working Forecast List.xlsx",
                           active_rows, complete_rows, with_inv=with_inv)
        src = tmp_path / "SSPO PRICING.xlsx"
        _pricing_workbook(src, pricing_rows)
        out = tmp_path / "out"
        exporter = exporter if exporter is not None else FakeExporter()
        logs = []
        result = mod.split_workbook(
            src, out, {"forecast_dir": str(fdir)}, logs.append,
            date_range=(date(2026, 9, 1), date(2026, 9, 7)),
            pdf_exporter=exporter)
        return result, out, exporter, logs
    return _run


def _supplement(out: Path, batch, nest, mod):
    wb = load_workbook(out / f"{batch} {nest} Invoicing Docs"
                       / f"{batch} {nest} Pricing Back Up.xlsx")
    return wb[mod.INV_SHEET]


# ---------------------------------------------------------------------------------
# Forecast lookup
# ---------------------------------------------------------------------------------
def test_lookup_carries_invoice_and_ship_date_from_both_sheets(mod, tmp_path):
    p = tmp_path / "wf.xlsx"
    _forecast_workbook(
        p,
        active_rows=[("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), "55501")],
        complete_rows=[("1000129724", 11, "V094", "503891", datetime(2026, 8, 20), "55420"),
                       ("1000129724", 14, "S038", "P08348", datetime(2025, 1, 1), "OLD")])
    m = mod._read_po_map(p, lambda *_: None, None)
    assert m[("S038", "P08348")] == mod.ForecastRow(
        "1000129724", 14, "55501", datetime(2026, 9, 3))   # active sheet wins
    assert m[("V094", "503891")].invoice == "55420"
    assert m[("V094", "503891")].ship_date == datetime(2026, 8, 20)


def test_lookup_without_the_column_warns_once_and_leaves_blank(mod, tmp_path):
    p = tmp_path / "wf.xlsx"
    _forecast_workbook(
        p, [("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), None)], [],
        with_inv=False)
    logs = []
    m = mod._read_po_map(p, logs.append, None)
    row = m[("S038", "P08348")]
    assert row.invoice == "" and row.po == "1000129724"
    assert sum(mod.INV_HEADER in l for l in logs) == 2    # one warning per sheet


# ---------------------------------------------------------------------------------
# Supplement + PDF
# ---------------------------------------------------------------------------------
def test_supplement_gets_invoice_and_ship_date_and_a_pdf(mod, run_split):
    result, out, exporter, _ = run_split(
        active_rows=[("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), "55501")],
        complete_rows=[],
        pricing_rows=[("S038", "P08348", 26308, 1234.5), ("S038", "P08348", 26309, 10)])
    ws = _supplement(out, "S038", "P08348", mod)
    assert ws["F2"].value == "Invoice #" and ws["G2"].value == "55501"
    assert ws["F3"].value == "Invoice Date:" and ws["G3"].value == datetime(2026, 9, 3)
    assert ws["G3"].number_format == "mm-dd-yy"

    docs = out / "S038 P08348 Invoicing Docs"
    assert exporter.calls == [(docs / "S038 P08348 Pricing Back Up.xlsx",
                               docs / "ASA Invoice No. 55501 Supplement.pdf")]
    assert (docs / "ASA Invoice No. 55501 Supplement.pdf").exists()
    assert result.pdfs == ["S038 P08348 Invoicing Docs\\ASA Invoice No. 55501 Supplement.pdf"]
    assert result.missing_invoice == [] and result.missing_po == []
    assert result.pdf_error is None and exporter.closed


def test_nest_without_invoice_number_gets_no_pdf_and_is_named(mod, run_split):
    result, out, exporter, logs = run_split(
        active_rows=[("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), None),
                     ("1000129724", 15, "S038", "P08356", datetime(2026, 9, 3), "55502")],
        complete_rows=[],
        pricing_rows=[("S038", "P08348", 1, 5), ("S038", "P08356", 2, 6),
                      ("V999", "504999", 3, 7)])      # V999: not on the forecast
    blank = _supplement(out, "S038", "P08348", mod)
    assert blank["G2"].value is None and blank["G3"].value == datetime(2026, 9, 3)
    assert blank["A" + str(mod.INV_DATA_START)].value == "1000129724"   # PO still filled

    assert [c[1].name for c in exporter.calls] == ["ASA Invoice No. 55502 Supplement.pdf"]
    assert result.missing_invoice == ["S038 P08348"]
    assert result.missing_po == ["V999 504999"]         # no forecast row -> not "missing invoice"
    assert any("P08348 has no PS/Inv" in l for l in logs)
    # Every workbook is still written regardless.
    assert len(result.written) == 3


def test_excel_unavailable_still_writes_workbooks(mod, run_split):
    exporter = FakeExporter(fail=True, error="Excel automation (pywin32) is not available")
    result, out, exporter, _ = run_split(
        active_rows=[("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), "55501")],
        complete_rows=[], pricing_rows=[("S038", "P08348", 1, 5)], exporter=exporter)
    assert result.written and result.pdfs == []
    assert result.pdf_error == "Excel automation (pywin32) is not available"
    assert _supplement(out, "S038", "P08348", mod)["G2"].value == "55501"


def test_logo_stays_above_the_header_band(mod, run_split):
    """2026-09-17: the printed PDF had the logo box overlapping the blue header
    band. The six title rows are pinned and the picture sized inside them."""
    _, out, _, _ = run_split(
        active_rows=[("1000129724", 14, "S038", "P08348", datetime(2026, 9, 3), "55501")],
        complete_rows=[], pricing_rows=[("S038", "P08348", 1, 5)])
    ws = _supplement(out, "S038", "P08348", mod)
    title_pt = sum(ws.row_dimensions[r].height for r in range(1, mod.INV_HEADER_ROW))
    assert title_pt == 6 * mod.TITLE_ROW_HEIGHT_PT
    assert mod.LOGO_PRINT_SIZE[1] <= title_pt * 96 / 72 - mod.LOGO_SEAM_PX
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1
    # A reloaded image reports its natural pixel size; the placed size is the
    # anchor extent, in EMU (9525 per px).
    anchor = ws._images[0].anchor
    assert (anchor.ext.width, anchor.ext.height) == tuple(px * 9525 for px in mod.LOGO_PRINT_SIZE)
    # Nudged 1 pt right and 1 pt down from A1 (Anthony's pick, 2026-09-17).
    assert (anchor._from.col, anchor._from.row) == (0, 0)
    assert (anchor._from.colOff, anchor._from.rowOff) == (12700, 12700)


def test_pdf_name_is_invoicings_own_and_filename_safe(mod):
    assert mod._supplement_pdf_name(" 55501 ") == "ASA Invoice No. 55501 Supplement.pdf"
    assert "/" not in mod._supplement_pdf_name("PS/55501")


def test_real_exporter_reports_missing_excel_without_raising(mod, monkeypatch, tmp_path):
    """The COM path degrades to a recorded error, never an exception, when
    pywin32 can't be imported - the workbooks must still be written."""
    import builtins
    real_import = builtins.__import__

    def no_win32(name, *a, **k):
        if name in ("pythoncom", "win32com.client", "win32com"):
            raise ImportError(name)
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", no_win32)
    logs = []
    ex = mod._SupplementPdfExporter(logs.append)
    assert ex.export(tmp_path / "a.xlsx", tmp_path / "a.pdf") is False
    assert ex.export(tmp_path / "b.xlsx", tmp_path / "b.pdf") is False
    assert "pywin32" in ex.error
    assert sum("no supplement PDFs" in l for l in logs) == 1      # warned once
    ex.close()
