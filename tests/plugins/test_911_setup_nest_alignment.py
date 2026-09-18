"""911 Setup v2.2.1 - the whole NEST tab is centered (floor feedback 2026-09-18:
"the top portion is centered and then further down it's left aligned").

Both SACO templates center only a patch of the sheet (A-E to row 25, the
batch-list columns F-K on the first few rows). The app now centers the sheet
itself after filling it, so the result never depends on how far the template's
formatting was dragged.
"""
import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment

PLUGIN = Path(__file__).resolve().parents[2] / "plugins" / "911_setup" / "run.py"


@pytest.fixture(scope="module")
def mod():
    spec = importlib.util.spec_from_file_location("sp_911_setup_align", PLUGIN)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _template_like_sheet(parts: int):
    """A NEST sheet formatted the way the real templates are: a centered patch
    at the top, nothing below it, with the app's pasted rows running past it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NEST"
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:C1")
    ws["A1"] = "NEST DATA"
    ws["A1"].alignment = center
    for c, h in enumerate(["FORECAST PO", "FORECAST PO LINE", "BATCH", "MIL SPEC",
                           "MATERIAL TYPE", "WORK ORDER", "DYPN", "MATERIAL",
                           "DYPN QTY", "NEST", "SCOPE OF WORK"], 1):
        ws.cell(3, c, h).alignment = center
    for r in range(4, 9):                       # the template's centered patch
        for c in range(1, 6):
            ws.cell(r, c).alignment = center
    for i in range(parts):                      # what the app pastes
        r = 4 + i
        for c in range(1, 12):
            ws.cell(r, c, f"v{r}-{c}")
    ws["G6"].alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
    ws["P9"] = "side note"                      # a value outside the table
    return wb, ws


def test_every_pasted_row_is_centered_past_the_templates_patch(mod):
    wb, ws = _template_like_sheet(parts=30)
    assert ws.cell(20, 7).alignment.horizontal != "center"      # the bug, before
    changed = mod._center_nest_sheet(ws, 30)
    assert changed > 0
    for r in range(3, 34):
        for c in range(1, 12):
            assert ws.cell(r, c).alignment.horizontal == "center", (r, c)
    assert ws["P9"].alignment.horizontal == "center"            # values anywhere on the tab
    assert ws["A1"].alignment.horizontal == "center"


def test_only_horizontal_changes_and_merged_cells_are_skipped(mod):
    wb, ws = _template_like_sheet(parts=5)
    mod._center_nest_sheet(ws, 5)
    g6 = ws["G6"].alignment
    assert (g6.horizontal, g6.wrap_text, g6.vertical) == ("center", True, "top")
    # B1/C1 are merged followers: no style of their own, must not raise or be touched
    assert type(ws["B1"]).__name__ == "MergedCell"


def test_blank_table_cells_are_centered_for_hand_entry_but_empty_space_is_left_alone(mod):
    wb, ws = _template_like_sheet(parts=3)
    ws.cell(5, 9).value = None                                   # a blank qty inside the table
    mod._center_nest_sheet(ws, 3)
    assert ws.cell(5, 9).alignment.horizontal == "center"
    assert ws.cell(40, 3).alignment.horizontal != "center"      # far below the data: untouched
    # second pass is a no-op
    assert mod._center_nest_sheet(ws, 3) == 0
